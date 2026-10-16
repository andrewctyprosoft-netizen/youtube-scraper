"""Build one Word document holding many transcripts, a heading per video."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

GREY = RGBColor(0x6B, 0x6B, 0x6B)
LINK = RGBColor(0x1A, 0x5F, 0xB4)


@dataclass
class DocSection:
    """One video (or file) in the combined document."""

    title: str
    blocks: list[str]
    url: str = ""
    info: list[str] = field(default_factory=list)  # e.g. channel, date, length


@dataclass
class DocOptions:
    """Layout switches for :func:`build_docx`."""

    doc_title: str = "Transcripts"
    contents: bool = True
    page_breaks: bool = True
    show_info: bool = True
    numbered: bool = False
    font: str = "Calibri"
    size: int = 11
    labels: dict[str, str] = field(
        default_factory=lambda: {
            "contents": "Contents",
            "created": "Created",
            "videos": "videos",
            "empty": "(No transcript text)",
        }
    )


def _set_font(document: Document, font: str, size: int) -> None:
    style = document.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    for name in ("Title", "Heading 1"):
        hstyle = document.styles[name]
        hstyle.font.name = font
        hrpr = hstyle.element.get_or_add_rPr()
        hfonts = hrpr.find(qn("w:rFonts"))
        if hfonts is None:
            hfonts = OxmlElement("w:rFonts")
            hrpr.append(hfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            hfonts.set(qn(attr), font)
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            hfonts.attrib.pop(qn(theme_attr), None)


def _add_bookmark(paragraph, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def _link_run(text: str, color: RGBColor) -> OxmlElement:
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), str(color))
    rpr.append(col)
    und = OxmlElement("w:u")
    und.set(qn("w:val"), "single")
    rpr.append(und)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    return run


def add_external_link(paragraph, url: str, text: str) -> None:
    """Append a clickable web link to ``paragraph``."""
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    link.append(_link_run(text, LINK))
    paragraph._p.append(link)


def add_internal_link(paragraph, anchor: str, text: str) -> None:
    """Append a link that jumps to bookmark ``anchor`` inside the document."""
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.append(_link_run(text, LINK))
    paragraph._p.append(link)


def build_docx(sections: list[DocSection], path: Path, options: DocOptions | None = None) -> Path:
    """Write ``sections`` to ``path`` (.docx) and return the path."""
    opts = options or DocOptions()
    labels = opts.labels
    document = Document()
    _set_font(document, opts.font, opts.size)

    document.add_paragraph(opts.doc_title, style="Title")
    meta = document.add_paragraph()
    run = meta.add_run(
        f"{labels['created']}: {datetime.now():%Y-%m-%d %H:%M} · {len(sections)} {labels['videos']}"
    )
    run.font.color.rgb = GREY
    run.font.size = Pt(9)

    if opts.contents and len(sections) > 1:
        document.add_paragraph(labels["contents"], style="Heading 2")
        for n, section in enumerate(sections, 1):
            item = document.add_paragraph()
            item.paragraph_format.space_after = Pt(2)
            add_internal_link(item, f"v{n}", f"{n}. {section.title}")
        if opts.page_breaks:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for n, section in enumerate(sections, 1):
        if n > 1 and opts.page_breaks:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        title = f"{n}. {section.title}" if opts.numbered else section.title
        heading = document.add_heading(title, level=1)
        _add_bookmark(heading, f"v{n}", n)

        if opts.show_info and (section.info or section.url):
            info = document.add_paragraph()
            info.paragraph_format.space_after = Pt(10)
            if section.info:
                r = info.add_run(" · ".join(section.info) + ("  ·  " if section.url else ""))
                r.font.color.rgb = GREY
                r.font.size = Pt(9)
            if section.url:
                add_external_link(info, section.url, section.url)

        if not section.blocks:
            r = document.add_paragraph().add_run(labels["empty"])
            r.italic = True
            r.font.color.rgb = GREY
        for block in section.blocks:
            para = document.add_paragraph(block)
            para.paragraph_format.space_after = Pt(6)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path
