import os
import re
from docx import Document

def natural_key(filename):
    """
    Key function for natural sorting of filenames.
    Splits the filename into numeric and non-numeric parts.
    """
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', filename)]

def combine_files(file_extension, file_type_name):
    """
    Generic function to combine files of a specific type into a Word document.
    """
    # Get folder path from user
    folder_path = input(f"Enter the path to the folder containing {file_type_name} files: ").strip()
    
    # Validate folder exists
    if not os.path.isdir(folder_path):
        print("Invalid folder path. Please provide a valid directory.")
        return
    
    # Get sorting option
    print("Sorting options:")
    print("1: Sort by filename (natural order: alphabetical with numerical awareness)")
    print("2: Sort by creation date and time")
    option = input("Enter option (1 or 2): ").strip()
    
    # Get output file name
    output_file = input("Enter the output Word file name (e.g., combined.docx): ").strip()
    if not output_file.endswith('.docx'):
        output_file += '.docx'
    
    # Find all files with the specified extension
    files = [f for f in os.listdir(folder_path) if f.lower().endswith(file_extension)]
    
    if not files:
        print(f"No {file_type_name} files found in the folder.")
        return
    
    # Sort based on option
    if option == '1':
        # Natural sort for filenames (handles numbers like 1, 2, 10)
        files.sort(key=natural_key)
    elif option == '2':
        # Sort by creation time (earliest first)
        files.sort(key=lambda f: os.path.getctime(os.path.join(folder_path, f)))
    else:
        print("Invalid option. Exiting.")
        return
    
    # Create Word document
    doc = Document()
    
    for file in files:
        # Add filename as heading
        doc.add_heading(file, level=1)
        
        # Read file content
        file_path = os.path.join(folder_path, file)
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
        except UnicodeDecodeError:
            # Try with different encoding if utf-8 fails
            try:
                with open(file_path, 'r', encoding='latin-1') as f:
                    content = f.read()
            except Exception as e:
                print(f"Error reading {file}: {e}")
                continue
        
        # Add content as a single paragraph (preserves original formatting as text)
        doc.add_paragraph(content)
        
        # Add a page break between files (optional, comment out if not needed)
        doc.add_page_break()
    
    # Save the document
    doc.save(output_file)
    print(f"Combined Word file saved as: {output_file}")

def main():
    print("=" * 50)
    print("         FILE COMBINER TOOL")
    print("=" * 50)
    print("Choose an option:")
    print("1: Combine TXT files into one Word document")
    print("2: Combine SRT files into one Word document")
    print("3: Exit")
    print("=" * 50)
    
    choice = input("Enter your choice (1, 2, or 3): ").strip()
    
    if choice == '1':
        print("\n--- Combining TXT Files ---")
        combine_files('.txt', 'TXT')
    elif choice == '2':
        print("\n--- Combining SRT Files ---")
        combine_files('.srt', 'SRT')
    elif choice == '3':
        print("Goodbye!")
        return
    else:
        print("Invalid choice. Please run the program again and select 1, 2, or 3.")
        return
    
    # Ask if user wants to perform another operation
    print("\n" + "=" * 50)
    another = input("Would you like to combine more files? (y/n): ").strip().lower()
    if another in ['y', 'yes']:
        print("\n")
        main()  # Recursive call to show menu again
    else:
        print("Thank you for using File Combiner Tool!")

if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\nProgram interrupted by user. Goodbye!")
    except Exception as e:
        print(f"\nAn unexpected error occurred: {e}")
    finally:
        # Keep the command window open
        input("\nPress Enter to exit...")
