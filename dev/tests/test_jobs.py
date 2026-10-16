"""Download jobs with a fake yt-dlp and fake caption server (no network)."""
from __future__ import annotations

import csv
import json
import re
import threading
import urllib.error
from pathlib import Path

import pytest
from docx import Document

from conftest import ROLLING_VTT, SRT_FROM_FFMPEG
from ytscraper import jobs
from ytscraper.engine import Cancelled, Engine, Entry, Tools
from ytscraper.jobs import (
    DownloadJob,
    FileJobOptions,
    JobOptions,
    choose_caption,
    combine_files,
    pick_fallback_language,
    safe_filename,
    title_from_filename,
    track_url,
)

EN_VTT = "WEBVTT\n\n00:00:01.000 --> 00:00:02.000\nHello world\n\n00:00:02.000 --> 00:00:03.000\nThis is 2026\n"

# subs = human captions, auto = YouTube auto captions (xx-orig = spoken language,
# others = machine translations). caption_429: how many times a caption GET answers 429
# ("always" = every time). extract_429: how many extraction attempts answer 429.
WORLD = {
    "aaaaaaaaaa1": {"title": "Bản tin 15/10 | Hải Phòng", "language": "vi", "subs": {},
                    "auto": {"vi": ROLLING_VTT, "vi-orig": ROLLING_VTT, "en": EN_VTT}},
    "bbbbbbbbbb2": {"title": "English talk: part 1", "language": "en", "subs": {"en": EN_VTT},
                    "auto": {"en-orig": EN_VTT, "vi": EN_VTT}},
    "ccccccccccc": {"error": "ERROR: [youtube] ccccccccccc: Private video. Sign in if you've been granted access"},
    "dddddddddd4": {"title": "Rate limited once", "language": "vi", "subs": {"vi": EN_VTT}, "auto": {},
                    "caption_429": {"vi": 1}},
    "eeeeeeeeee5": {"title": "No captions at all", "language": "", "subs": {}, "auto": {}},
}

# The real case from the AI Search channel (Oct 2026): plain "en" is a translated
# track that answers 429 every time; "en-orig" is the real auto-caption track.
AI_SEARCH = {
    "lHmZoRHMZyM": {"title": "Gemini 4, GPT 6.1", "language": "en-US", "subs": {},
                    "auto": {"en": EN_VTT, "en-orig": EN_VTT, "vi": EN_VTT},
                    "caption_429": {"en": "always", "vi": "always"}},
}


class FakeYtDlp:
    """Imitates the yt-dlp command line and YouTube's caption server."""

    def __init__(self, world: dict):
        self.world = json.loads(json.dumps(world))
        self.calls: list[list[str]] = []
        self.gets: list[str] = []

    # --- caption server (Engine.http_get)
    def get(self, url, timeout=30):
        self.gets.append(url)
        vid, code = url.removeprefix("fake://").split("/")
        video = self.world[vid]
        left = video.get("caption_429", {}).get(code, 0)
        if left == "always" or left > 0:
            if left != "always":
                video["caption_429"][code] -= 1
            raise urllib.error.HTTPError(url, 429, "Too Many Requests", {}, None)
        text = video["subs"].get(code) or video["auto"].get(code)
        return text.encode()

    def info(self, vid, url):
        video = self.world[vid]
        track = lambda code: [{"ext": "vtt", "url": f"fake://{vid}/{code}"}]  # noqa: E731
        return {"id": vid, "title": video["title"], "channel": "Chan", "upload_date": "20261015",
                "duration": 75, "view_count": 10, "webpage_url": url, "language": video["language"],
                "subtitles": {k: track(k) for k in video["subs"]},
                "automatic_captions": {k: track(k) for k in video["auto"]},
                "thumbnails": [{"url": f"fake://{vid}/thumb.jpg", "width": 1280}]}

    # --- yt-dlp command line (Engine.runner)
    def __call__(self, cmd, on_line=None, cancel=None):
        self.calls.append(cmd)
        args = cmd[cmd.index("--ignore-config"):]
        if "-j" in args:
            return self._batch(args, on_line)
        url = args[args.index("--") + 1] if "--" in args else ""
        vid = url.split("v=")[-1] if url else ""
        if "--load-info-json" in args:
            vid = json.loads(Path(args[args.index("--load-info-json") + 1]).read_text(encoding="utf-8"))["id"]
        video = self.world[vid]
        if "error" in video:
            return 1, video["error"]
        homes = [args[i + 1] for i, a in enumerate(args) if a == "-P" and ":" not in args[i + 1][2:]]
        home = Path(homes[0])
        home.mkdir(parents=True, exist_ok=True)
        lines = []
        if "--skip-download" in args and "--sub-langs" in args:
            patterns = args[args.index("--sub-langs") + 1].split(",")
            for lang, text in {**video["auto"], **video["subs"]}.items():
                if any(re.fullmatch(p, lang) for p in patterns):
                    (home / f"{vid}.{lang}.vtt").write_text(text, encoding="utf-8")
        elif "--skip-download" not in args:
            template = args[args.index("-o") + 1]
            ext = "mp3" if "-x" in args else "mp4"
            name = template.replace("%(title).150B", video["title"].replace("/", "-").replace(":", " -")).replace("%(ext)s", ext)
            target = home / name
            target.write_bytes(b"fake")
            for pct in ("10.0", "55.5", "100.0"):
                line = f"[download]  {pct}% of 1.00MiB"
                lines.append(line)
                if on_line:
                    on_line(line)
            lines.append(str(target))
        return 0, "\n".join(lines)

    def _batch(self, args, on_line):
        urls = Path(args[args.index("-a") + 1]).read_text(encoding="utf-8").split()
        code = 0
        for url in urls:
            vid = url.split("v=")[-1]
            video = self.world[vid]
            on_line(f"[youtube] {vid}: Downloading webpage")
            if "error" in video:
                on_line(video["error"])
                code = 1
            elif video.get("extract_429", 0) > 0:
                video["extract_429"] -= 1
                on_line(f"ERROR: [youtube] {vid}: Unable to download API page: HTTP Error 429: Too Many Requests")
                code = 1
            elif video.get("degraded", 0) > 0:
                video["degraded"] -= 1
                info = {**self.info(vid, url), "language": None, "subtitles": {}, "automatic_captions": {}}
                on_line(json.dumps(info))
            else:
                on_line(json.dumps(self.info(vid, url)))
        return code, ""


def make_engine(fake, tmp_path):
    tools = Tools(ytdlp=["yt-dlp"], deno=None, node=None, ffmpeg_dir=tmp_path, frozen_ytdlp=True)
    engine = Engine(tools, runner=fake)
    engine.http_get = fake.get
    return engine


@pytest.fixture(autouse=True)
def no_waiting(monkeypatch):
    monkeypatch.setattr(DownloadJob, "BACKOFF", (0, 0, 0))
    monkeypatch.setattr(DownloadJob, "PASS_WAIT", (0, 0))
    monkeypatch.setattr(DownloadJob, "CAPTION_GAP", 0)


@pytest.fixture
def fake():
    return FakeYtDlp(WORLD)


@pytest.fixture
def engine(fake, tmp_path):
    return make_engine(fake, tmp_path)


def entries(*ids):
    return [Entry(i, i, f"https://www.youtube.com/watch?v={i}") for i in ids]


class Recorder:
    def __init__(self):
        self.events = []

    def __call__(self, kind, **data):
        self.events.append((kind, data))

    def of(self, kind):
        return [d for k, d in self.events if k == kind]


def test_transcripts_end_to_end(engine, fake, tmp_path):
    out = tmp_path / "out"
    opts = JobOptions(out_dir=out, language="vi", save_txt=True, save_srt=True, save_combined_txt=True,
                      delay=0, output_name="My doc", numbered=True)
    rec = Recorder()
    results = DownloadJob(engine, entries(*WORLD), opts, rec).run()
    status = {r.entry.id: (r.status, r.language, r.error) for r in results}
    assert status["aaaaaaaaaa1"] == ("done", "vi-orig", "")  # real captions, not the "vi" translation
    assert status["bbbbbbbbbb2"] == ("done", "en", "")  # not spoken in vi: the video's own language
    assert status["ccccccccccc"] == ("error", "", "private")
    assert status["dddddddddd4"] == ("done", "vi", "")  # caption 429 once, then back-off and success
    assert status["eeeeeeeeee5"][0] == "no_captions"
    assert not any(u.endswith("/vi") and "aaaaaaaaaa1" in u for u in fake.gets)

    doc = Document(str(out / "My doc.docx"))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Bản tin 15/10 | Hải Phòng", "English talk: part 1", "ccccccccccc",
                        "Rate limited once", "No captions at all"]  # list order, not finish order
    body = "\n".join(p.text for p in doc.paragraphs)
    assert body.count("Kính chào quý vị") == 1
    assert "[âm nhạc]" not in body
    assert len(doc.element.body.findall(".//{*}hyperlink")) >= 5

    txt = (out / "Transcripts" / "001 - Bản tin 15-10 - Hải Phòng.txt").read_text(encoding="utf-8")
    assert txt.startswith("Bản tin 15/10 | Hải Phòng\nhttps://www.youtube.com/watch?v=aaaaaaaaaa1")
    assert (out / "Subtitles" / "002 - English talk - part 1.en.srt").exists()
    assert (out / "My doc.txt").read_text(encoding="utf-8").count("=" * 60) == 4

    with open(out / "My doc - list.csv", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    assert rows[0][:3] == ["No", "Title", "URL"]
    assert rows[3][9] == "error: private"
    assert rows[1][4] == "2026-10-15" and rows[1][5] == "1:15"

    done = rec.of("done")[0]
    assert done["counts"] == {"done": 3, "error": 1, "no_captions": 1}
    assert any(d.get("text") == "rate_limited_wait" for d in rec.of("progress"))
    assert rec.of("progress")[-2]["done"] == 5
    assert not list((out / ".cache").glob("*.info.json"))
    # one yt-dlp process per worker for the whole list, no per-video start-up
    assert len(fake.calls) == 3 and all("-j" in c for c in fake.calls)  # 2 workers + 1 retry of the caption-less video


def test_ai_search_case_uses_en_orig_and_never_blocks(tmp_path):
    fake = FakeYtDlp(AI_SEARCH)
    opts = JobOptions(out_dir=tmp_path, language="en", delay=0, save_csv=False)
    results = DownloadJob(make_engine(fake, tmp_path), entries("lHmZoRHMZyM"), opts, Recorder()).run()
    assert (results[0].status, results[0].language) == ("done", "en-orig")
    assert fake.gets == ["fake://lHmZoRHMZyM/en-orig"]


def test_translation_only_when_asked_and_backs_off(tmp_path):
    fake = FakeYtDlp(AI_SEARCH)
    rec = Recorder()
    opts = JobOptions(out_dir=tmp_path, language="vi", fallback_original=False, delay=0, save_csv=False)
    results = DownloadJob(make_engine(fake, tmp_path), entries("lHmZoRHMZyM"), opts, rec).run()
    assert (results[0].status, results[0].error) == ("error", "rate_limited")
    # 3 passes x (1 try + 3 back-offs) on the translated track, then yt-dlp's own attempt is not used
    assert fake.gets.count("fake://lHmZoRHMZyM/vi") == 12
    assert len(rec.of("row")[-1:]) == 1 and rec.of("row")[-1]["status"] == "error"


def test_extraction_429_is_retried_in_a_later_pass(tmp_path):
    world = {"bbbbbbbbbb2": {**WORLD["bbbbbbbbbb2"], "extract_429": 1}}
    fake = FakeYtDlp(world)
    rec = Recorder()
    opts = JobOptions(out_dir=tmp_path, language="en", delay=0, save_csv=False)
    results = DownloadJob(make_engine(fake, tmp_path), entries("bbbbbbbbbb2"), opts, rec).run()
    assert results[0].status == "done"
    assert sum("-j" in c for c in fake.calls) == 2
    assert any(d.get("text") == "retry_wait" for d in rec.of("progress"))


def test_timed_out_info_is_retried_not_reported_as_no_captions(tmp_path):
    # Real case: under network timeouts yt-dlp returned info without language or caption lists.
    world = {"bbbbbbbbbb2": {**WORLD["bbbbbbbbbb2"], "degraded": 1}, "eeeeeeeeee5": WORLD["eeeeeeeeee5"]}
    fake = FakeYtDlp(world)
    opts = JobOptions(out_dir=tmp_path, language="en", delay=0, save_csv=False)
    results = DownloadJob(make_engine(fake, tmp_path), entries(*world), opts, Recorder()).run()
    assert [(r.status, r.language) for r in results] == [("done", "en"), ("no_captions", "")]
    assert sum("-j" in c for c in fake.calls) == 2 + 2  # first pass + one retry pass (2 workers each)


def test_direct_download_refused_falls_back_to_ytdlp(engine, fake, tmp_path):
    def refuse(url, timeout=30):
        raise urllib.error.HTTPError(url, 403, "Forbidden", {}, None)

    engine.http_get = refuse
    opts = JobOptions(out_dir=tmp_path, language="en", delay=0, save_csv=False)
    results = DownloadJob(engine, entries("bbbbbbbbbb2"), opts, Recorder()).run()
    assert (results[0].status, results[0].language) == ("done", "en")
    assert any("--load-info-json" in c for c in fake.calls)


def test_second_run_uses_cache(engine, fake, tmp_path):
    out = tmp_path / "out"
    opts = JobOptions(out_dir=out, language="vi", delay=0, save_csv=False)
    DownloadJob(engine, entries("aaaaaaaaaa1", "bbbbbbbbbb2"), opts, Recorder()).run()
    calls, gets = len(fake.calls), len(fake.gets)
    results = DownloadJob(engine, entries("aaaaaaaaaa1", "bbbbbbbbbb2"), opts, Recorder()).run()
    assert (len(fake.calls), len(fake.gets)) == (calls, gets)  # nothing downloaded again
    assert [r.language for r in results] == ["vi-orig", "en"]
    assert all(r.from_cache for r in results)


def test_original_language_choice(engine, tmp_path):
    opts = JobOptions(out_dir=tmp_path, language="orig", delay=0)
    results = DownloadJob(engine, entries("aaaaaaaaaa1", "bbbbbbbbbb2"), opts, Recorder()).run()
    assert [r.language for r in results] == ["vi-orig", "en"]  # human captions beat auto ones


def test_no_fallback_uses_translation(engine, fake, tmp_path):
    opts = JobOptions(out_dir=tmp_path, language="vi", fallback_original=False, delay=0)
    results = DownloadJob(engine, entries("bbbbbbbbbb2"), opts, Recorder()).run()
    assert (results[0].status, results[0].language) == ("done", "vi")
    assert fake.gets == ["fake://bbbbbbbbbb2/vi"]


def test_media_downloads(engine, tmp_path):
    rec = Recorder()
    opts = JobOptions(out_dir=tmp_path, want_transcript=False, want_video=True, want_audio=True,
                      save_csv=False, delay=0)
    results = DownloadJob(engine, entries("bbbbbbbbbb2"), opts, rec).run()
    names = sorted(p.name for p in results[0].files)
    assert names == ["English talk - part 1.mp3", "English talk - part 1.mp4"]
    assert results[0].status == "done"
    assert any(d.get("detail") == "video 56%" for d in rec.of("row"))
    assert not any("%" in d["line"] for d in rec.of("log"))


def test_thumbnail_is_downloaded_directly(engine, fake, tmp_path):
    fake.get = lambda url, timeout=30: b"\xff\xd8\xff\xe0jpeg" if url.endswith("thumb.jpg") else b""
    engine.http_get = fake.get
    opts = JobOptions(out_dir=tmp_path, want_transcript=False, want_thumbnail=True, save_csv=False, delay=0)
    results = DownloadJob(engine, entries("bbbbbbbbbb2"), opts, Recorder()).run()
    assert [p.name for p in results[0].files] == ["English talk - part 1.jpg"]
    # a video without captions still gets its thumbnail when transcripts are wanted too
    opts = JobOptions(out_dir=tmp_path / "b", want_thumbnail=True, save_csv=False, delay=0)
    results = DownloadJob(engine, entries("eeeeeeeeee5"), opts, Recorder()).run()
    assert results[0].status == "no_captions"
    assert [p.suffix for p in results[0].files] == [".jpg"]


def test_cancel_stops_and_still_finishes(engine, tmp_path):
    cancel = threading.Event()
    rec = Recorder()

    def emit(kind, **data):
        rec(kind, **data)
        if kind == "row" and data["status"] not in ("working",):
            cancel.set()

    opts = JobOptions(out_dir=tmp_path, delay=0, workers=1)
    with pytest.raises(Cancelled):
        DownloadJob(engine, entries("aaaaaaaaaa1", "bbbbbbbbbb2"), opts, emit, cancel).run()
    done = rec.of("done")[0]
    assert done["cancelled"] is True
    assert done["counts"] == {"done": 1}  # the unfinished video is left out
    assert any(p.endswith(".docx") for p in done["outputs"])


def test_locked_word_file_gets_new_name(engine, tmp_path, monkeypatch):
    real = jobs.build_docx

    def flaky(sections, path, options):
        if path.name == "Transcripts.docx":
            raise PermissionError(path)
        return real(sections, path, options)

    monkeypatch.setattr(jobs, "build_docx", flaky)
    opts = JobOptions(out_dir=tmp_path, delay=0, output_name="Transcripts", save_csv=False)
    rec = Recorder()
    DownloadJob(engine, entries("aaaaaaaaaa1"), opts, rec).run()
    assert rec.of("done")[0]["outputs"][0].endswith("Transcripts (2).docx")


def test_title_sort_key_vietnamese_numbers_symbols():
    titles = ["Đà Nẵng", "ánh sáng", "Bản tin", "Part 10", "Part 2", '"Quoted" title', "🔥 Hot news", "zebra"]
    ordered = [titles[i] for i in jobs.file_order(titles, "az")]
    assert ordered == ["ánh sáng", "Bản tin", "Đà Nẵng", "🔥 Hot news", "Part 2", "Part 10", '"Quoted" title', "zebra"]
    assert [titles[i] for i in jobs.file_order(titles, "za")] == list(reversed(ordered))
    assert jobs.file_order(titles, "list") == list(range(len(titles)))


@pytest.mark.parametrize("order,expected", [
    ("list", ["001 - Bản tin 15-10 - Hải Phòng", "002 - English talk - part 1", "003 - Rate limited once"]),
    ("az", ["001 - Bản tin 15-10 - Hải Phòng", "002 - English talk - part 1", "003 - Rate limited once"]),
    ("za", ["001 - Rate limited once", "002 - English talk - part 1", "003 - Bản tin 15-10 - Hải Phòng"]),
])
def test_order_numbers_files_and_documents(engine, tmp_path, order, expected):
    ids = ["dddddddddd4", "aaaaaaaaaa1", "bbbbbbbbbb2"]  # list order differs from A→Z
    titles = {i: WORLD[i]["title"] for i in ids}
    ents = [Entry(i, titles[i], f"https://www.youtube.com/watch?v={i}") for i in ids]
    if order == "list":
        expected = ["001 - Rate limited once", "002 - Bản tin 15-10 - Hải Phòng", "003 - English talk - part 1"]
    opts = JobOptions(out_dir=tmp_path, language="orig", delay=0, save_txt=True, numbered=True, order=order,
                      want_video=True, output_name="Doc")
    DownloadJob(engine, ents, opts, Recorder()).run()
    assert sorted(p.stem for p in (tmp_path / "Transcripts").iterdir()) == expected
    # video files carry exactly the same name as the transcript of the same video
    assert sorted(p.stem for p in (tmp_path / "Videos").iterdir()) == expected
    doc = Document(str(tmp_path / "Doc.docx"))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert [h.replace("/", "-").replace(" |", " -").replace(":", " -") for h in headings] == [e[6:] for e in expected]
    with open(tmp_path / "Doc - list.csv", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))[1:]
    assert [r[0] for r in rows] == ["1", "2", "3"]
    assert [r[1] for r in rows] == headings


def test_numbers_widen_for_big_lists(engine, tmp_path):
    job = DownloadJob(engine, [], JobOptions(out_dir=tmp_path, numbered=True), Recorder())
    job.rank, job.width = {0: 1233}, 4
    assert job._media_prefix(0) == "1234 - "


@pytest.mark.parametrize("meta,requested,kwargs,expected", [
    # spoken English, only auto captions: en-orig, never the translated "en"
    ({"language": "en-US", "subtitles": [], "automatic_captions": ["en", "en-orig", "vi"]}, "en", {}, ("en-orig", "auto")),
    # human captions win
    ({"language": "en", "subtitles": ["en-GB"], "automatic_captions": ["en", "en-orig"]}, "en", {}, ("en-GB", "manual")),
    # not spoken in Vietnamese: fall back to the video's own language
    ({"language": "en", "subtitles": [], "automatic_captions": ["en-orig", "vi"]}, "vi", {}, ("en-orig", "auto")),
    # ... unless the user wants the translation
    ({"language": "en", "subtitles": [], "automatic_captions": ["en-orig", "vi"]}, "vi",
     {"fallback_original": False}, ("vi", "translated")),
    # auto captions not allowed
    ({"language": "en", "subtitles": [], "automatic_captions": ["en-orig"]}, "en", {"allow_auto": False}, None),
    # older listing without -orig tracks
    ({"language": "en", "subtitles": [], "automatic_captions": ["en"]}, "en", {}, ("en", "auto")),
    ({"language": "vi", "subtitles": [], "automatic_captions": ["en-orig", "vi-orig"]}, "orig", {}, ("vi-orig", "auto")),
])
def test_choose_caption(meta, requested, kwargs, expected):
    assert choose_caption(meta, requested, **kwargs) == expected


def test_track_url_prefers_small_srt():
    tracks = [{"ext": "json3", "url": "j"}, {"ext": "vtt", "url": "v"}, {"ext": "srt", "url": "s"}]
    assert track_url({"automatic_captions": {"en-orig": tracks}}, "en-orig") == "s"
    assert track_url({"automatic_captions": {"en-orig": tracks[:2]}}, "en-orig") == "v"
    assert track_url({"automatic_captions": {"en-orig": tracks[:1]}}, "en-orig") is None  # unreadable only
    assert track_url({"automatic_captions": {"en-orig": tracks}}, "fr") is None


def test_caption_timeouts_are_retried_without_extra_ytdlp(engine, fake, tmp_path):
    real, fails = fake.get, [2]

    def flaky(url, timeout=30):
        if fails[0]:
            fails[0] -= 1
            raise urllib.error.URLError("timed out")
        return real(url, timeout)

    engine.http_get = flaky
    opts = JobOptions(out_dir=tmp_path, language="en", delay=0, save_csv=False)
    results = DownloadJob(engine, entries("bbbbbbbbbb2"), opts, Recorder()).run()
    assert results[0].status == "done"
    assert not any("--load-info-json" in c for c in fake.calls)


def test_caption_saved_with_matching_extension(engine, fake, tmp_path):
    srt = "1\n00:00:01,000 --> 00:00:02,000\nHello from SRT\n"
    fake.world["bbbbbbbbbb2"]["subs"]["en"] = srt
    opts = JobOptions(out_dir=tmp_path, language="en", delay=0, save_csv=False)
    results = DownloadJob(engine, entries("bbbbbbbbbb2"), opts, Recorder()).run()
    assert results[0].blocks == ["Hello from SRT"]
    assert (tmp_path / ".cache" / "bbbbbbbbbb2.en.srt").exists()


def test_combine_my_files(tmp_path):
    (tmp_path / "10 - Talk [dQw4w9WgXcQ].vi.vtt").write_text(ROLLING_VTT, encoding="utf-8")
    (tmp_path / "2 - Show.en.srt").write_text(SRT_FROM_FFMPEG, encoding="utf-8")
    (tmp_path / "notes.txt").write_text("line one\nline one\nline two\n", encoding="utf-16")
    paths = sorted(tmp_path.iterdir(), key=lambda p: jobs.natural_key(p.name))
    out = tmp_path / "out"
    opt = FileJobOptions(out_dir=out, save_txt=True, save_srt=True, output_name="All")
    rec = Recorder()
    written = combine_files(paths, opt, rec)
    doc = Document(str(out / "All.docx"))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["2 - Show", "10 - Talk", "notes"]
    assert "line one line two" in "\n".join(p.text for p in doc.paragraphs)
    assert (out / "2 - Show.clean.srt").exists()
    assert not (out / "notes.clean.srt").exists()
    assert len(written) == 6


@pytest.mark.parametrize("name,title", [
    ("012 - My talk [dQw4w9WgXcQ].vi-orig.srt", "012 - My talk"),
    ("Bản tin.vi.vtt", "Bản tin"),
    ("Lecture 3.txt", "Lecture 3"),
    ("Show.en-US.srt", "Show"),
])
def test_title_from_filename(name, title):
    assert title_from_filename(name) == title


def test_safe_filename():
    assert safe_filename("A/B: C? | D*") == "A-B - C - D"
    assert safe_filename("CON") == "_CON"
    assert safe_filename("x" * 300) == "x" * 150


def test_pick_fallback_language():
    assert pick_fallback_language({"language": "en", "subtitles": ["de", "en-GB"], "automatic_captions": ["en-orig"]}) == "en-GB"
    assert pick_fallback_language({"language": "", "subtitles": [], "automatic_captions": ["en", "ja-orig"]}) == "ja-orig"
    assert pick_fallback_language({"subtitles": ["live_chat"], "automatic_captions": []}) is None


def test_pick_fallback_ignores_ai_dub_tracks():
    meta = {"language": "vi", "subtitles": [], "automatic_captions": ["en", "en-orig", "vi", "vi-orig"]}
    assert pick_fallback_language(meta) == "vi-orig"
    meta = {"language": "", "subtitles": ["ar", "bg", "en-US", "es"], "automatic_captions": []}
    assert pick_fallback_language(meta) == "en-US"
    meta = {"language": "en", "subtitles": ["ar", "iw", "zh-TW"], "automatic_captions": ["en-orig"]}
    assert pick_fallback_language(meta) == "en-orig"


def test_stale_cache_languages_are_not_used(tmp_path):
    (tmp_path / "aaaaaaaaaa1.vi-orig.vtt").write_text(ROLLING_VTT, encoding="utf-8")
    assert jobs.find_subtitle(tmp_path, "aaaaaaaaaa1", "en") is None
    assert jobs.find_subtitle(tmp_path, "aaaaaaaaaa1", "vi").name == "aaaaaaaaaa1.vi-orig.vtt"
